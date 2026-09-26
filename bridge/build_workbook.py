"""銘柄セットから RSS 関数を敷いた Excel ブック（rss_bridge.xlsx）を生成する。

生成物の quotes シート:

  A列: code   B列: 現在値               C列: 出来高               ...
  2:   7203   =RssMarket(A2,"現在値")   =RssMarket(A2,"出来高")   ...
  3:   6501   ...

Windows でこのブックを Excel で開くと（マーケットスピードII にログイン済みなら）
RSS 関数が値を返し、bridge.py がその表を読んで backend に送る。

  python build_workbook.py --set-id 1
  python build_workbook.py --codes 7203,6501,9984 --out rss_bridge.xlsx

※ フィールド名（"現在値" 等）は楽天証券公式の RSS リファレンスで要確認。
   VBA を足したい場合は Excel で .xlsm として保存し直す。
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from sqlmodel import Session, select  # noqa: E402

from app.db import engine, init_db  # noqa: E402
from app.models import SymbolSet, SymbolSetItem  # noqa: E402

# bridge.py の FIELDS と必ず一致させること
FIELDS = ["現在値", "出来高", "前日比", "最良買気配値", "最良売気配値"]

# --probe 用。RssMarket の項目名候補を総当たりで並べ、Windows で開いてどれが
# 数値を返すか目視確認する（楽天証券 RSS リファレンス未確定のため）。
PROBE_FIELDS = [
    "銘柄名称", "現在値", "現在値時刻", "出来高", "売買代金",
    "始値", "高値", "安値", "前日終値", "前日比", "前日比率",
    "最良売気配値", "最良買気配値",
    "最良売気配値1", "最良買気配値1",
    "売気配値1", "買気配値1",
    "最良売気配数量", "最良買気配数量",
    "VWAP", "約定回数", "市場コード",
]

# --orders 用。RssStockOrder（国内株式・現物注文）の引数（20個、A〜T列）。
# bridge.py の ORDER_INPUT_COLS と対応関係を保つこと。
ORDER_ARG_HEADER = [
    "発注ID", "発注トリガー", "銘柄コード", "売買区分", "注文区分", "SOR区分",
    "注文数量", "価格区分", "注文価格", "執行条件", "注文期限", "口座区分",
    "逆指値条件価格", "逆指値条件区分", "逆指値価格区分", "逆指値価格",
    "セット注文区分", "セット注文価格", "セット注文執行条件", "セット注文期限",
]


def load_codes(set_id: int) -> list[str]:
    init_db()
    with Session(engine) as s:
        if not s.get(SymbolSet, set_id):
            raise SystemExit(f"symbol set id={set_id} が見つかりません")
        items = s.exec(
            select(SymbolSetItem).where(SymbolSetItem.set_id == set_id).order_by(SymbolSetItem.sort_order)
        ).all()
        return [it.symbol_code for it in items]


def build(codes: list[str], out_path: Path) -> None:
    if not codes:
        raise SystemExit("銘柄が空です")
    try:
        from openpyxl import Workbook
    except ImportError as e:
        raise SystemExit("openpyxl が必要です: pip install openpyxl") from e

    wb = Workbook()
    ws = wb.active
    ws.title = "quotes"
    ws.append(["code", *FIELDS])
    for i, code in enumerate(codes, start=2):
        ws.cell(row=i, column=1, value=str(code))
        for j, f in enumerate(FIELDS, start=2):
            ws.cell(row=i, column=j, value=f'=RssMarket(A{i},"{f}")')
    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 10
    for col in ("B", "C", "D", "E", "F"):
        ws.column_dimensions[col].width = 14

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    print(f"生成: {out_path}  ({len(codes)} 銘柄)")
    print("Windows の Excel で開くと RSS 関数が評価されます（マーケットスピードII ログイン必須）。")


def build_probe(code: str, out_path: Path) -> None:
    """1銘柄ぶん、項目名候補を縦に並べたブックを生成する。"""
    try:
        from openpyxl import Workbook
    except ImportError as e:
        raise SystemExit("openpyxl が必要です: pip install openpyxl") from e

    wb = Workbook()
    ws = wb.active
    ws.title = "probe"
    ws.append(["項目名", f'RssMarket("{code}", 項目名)'])
    for i, f in enumerate(PROBE_FIELDS, start=2):
        ws.cell(row=i, column=1, value=f)
        ws.cell(row=i, column=2, value=f'=RssMarket("{code}","{f}")')
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 22
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    print(f"生成: {out_path}  (probe: {code}, {len(PROBE_FIELDS)} 項目)")
    print("Excel で開き、B列に数値/文字が返る項目名をメモ → FIELDS を修正する。")


def build_orders(out_path: Path, n_rows: int = 300) -> None:
    """発注専用ブック（1回だけ作って使い回す。quotes ブックとは別ファイル）。

    quotes ブックは銘柄セットを変えるたびに作り直すが、orders ブックは
    発注中の状態を持つので自動では作り直さない・書き換えない設計にしてある。
    """
    try:
        from openpyxl import Workbook
        from openpyxl.utils import get_column_letter
    except ImportError as e:
        raise SystemExit("openpyxl が必要です: pip install openpyxl") from e

    n_args = len(ORDER_ARG_HEADER)  # 20 (A〜T列)
    status_col = n_args + 1  # U列(21列目)

    wb = Workbook()
    ws = wb.active
    ws.title = "orders"
    ws.append([*ORDER_ARG_HEADER, "ステータス"])
    for row in range(2, n_rows + 2):
        cell_refs = ",".join(f"{get_column_letter(c)}{row}" for c in range(1, n_args + 1))
        ws.cell(row=row, column=status_col, value=f"=RssStockOrder({cell_refs})")
    for c in range(1, n_args + 1):
        ws.column_dimensions[get_column_letter(c)].width = 11
    ws.column_dimensions[get_column_letter(status_col)].width = 45
    ws.freeze_panes = "A2"

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    print(f"生成: {out_path}  (発注用シート、{n_rows}行ぶんの空き)")
    print("このファイルは1回作って Excel で開いたままにする（quotes ブックと違い自動では作り直さない）。")
    print("※ RssStockOrder の実際の発注は未検証。まず MarketSpeed II 側の「発注機能」を"
          "OFFのままテストし、「発注ロック中」が正しく検出できることを確認してから有効化すること。")


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--set-id", type=int)
    ap.add_argument("--codes", help="カンマ区切り 例: 7203,6501")
    ap.add_argument("--probe", help="項目名の実地確認用ブックを作る（1銘柄コード）")
    ap.add_argument("--orders", action="store_true", help="発注用ブック（rss_orders.xlsx）を作る")
    ap.add_argument("--out", default=str(Path(__file__).parent / "rss_bridge.xlsx"))
    args = ap.parse_args()

    if args.orders:
        default_out = Path(__file__).parent / "rss_orders.xlsx"
        out = Path(args.out) if args.out != str(Path(__file__).parent / "rss_bridge.xlsx") else default_out
        build_orders(out)
        return

    if args.probe:
        default_out = Path(__file__).parent / "rss_probe.xlsx"
        out = Path(args.out) if args.out != str(Path(__file__).parent / "rss_bridge.xlsx") else default_out
        build_probe(args.probe.strip(), out)
        return

    if args.codes:
        codes = [c.strip() for c in args.codes.split(",") if c.strip()]
    elif args.set_id:
        codes = load_codes(args.set_id)
    else:
        raise SystemExit("--set-id か --codes か --probe か --orders を指定してください")
    build(codes, Path(args.out))


if __name__ == "__main__":
    main()
